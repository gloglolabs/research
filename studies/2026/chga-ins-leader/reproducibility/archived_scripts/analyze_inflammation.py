"""Reanalyse published Carré2025 immunopeptidomics for a canonical INS SP1-24 replacement.
Run with bundled Python containing openpyxl, numpy and matplotlib.
No fitting or imputation. Condition mean/median intensities are source values,
not replicate measurements. HLA assignments are published predictions.
"""
from pathlib import Path
import csv, json, hashlib, math, re
from collections import Counter
from statistics import median
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
REF=json.loads((ROOT.parents[1]/'evidence/uniprot_reference_sequences.json').read_text())
SRC='https://www.nature.com/articles/s41467-025-55908-9'

def sheet(name, tab=None):
    wb=load_workbook(ROOT/name,read_only=True,data_only=True)
    return list(wb[tab or wb.sheetnames[0]].values)

def number(v):
    if v is None: return None
    try: return float(v)
    except (TypeError,ValueError): return None

def write_csv(name, rows):
    if not rows:return
    with (ROOT/name).open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

# Use each all-binders column only once, excluding duplicate HLA subset blocks.
quant={}
for rownum,row in enumerate(sheet('41467_2025_55908_MOESM14_ESM.xlsx','Fig 4d-e')[5:],6):
    for i in (1,13):
        if not isinstance(row[i],str): continue
        seq=row[i].strip(); vals=(number(row[i+1]),number(row[i+2]))
        assert seq not in quant or quant[seq]['values']==vals, seq
        quant[seq]={'values':vals,'source_cell':f"Fig 4d-e!{'B' if i==1 else 'N'}{rownum}"}
replication={}
for rownum,row in enumerate(sheet('41467_2025_55908_MOESM14_ESM.xlsx','Fig S3c')[3:],4):
    if isinstance(row[1],str):
        replication[row[1]]={'values':tuple(number(x) for x in row[2:7]),'source_cell':f'Fig S3c!B{rownum}'}

all_rows=[]
for rownum,r in enumerate(sheet('41467_2025_55908_MOESM4_ESM.xlsx')[1:],2):
    seq,presence,basal,ifna,length,uniprot,position,gene,synonyms,description,granule,hla,rank=r[:13]
    start,end=map(int,str(position).split('-')) if re.fullmatch(r'\d+-\d+',str(position)) else (None,None)
    gene_ref=next((g for g,ref in REF.items() if ref['accession'] in str(uniprot).split(';')),None)
    mapping=(REF[gene_ref]['sequence'][start-1:end]==seq) if gene_ref and start and end else None
    is_ins='P01308' in str(uniprot).split(';')
    if is_ins:
        segment='INS signal 1-24' if end<=24 else ('INS retained 25-110' if start>=25 else 'INS boundary overlap')
    elif gene_ref=='CHGA' and start and end<=18: segment='CHGA signal 1-18'
    else:segment='Other region'
    x=quant.get(seq,{});b,a=x.get('values',(None,None))
    fc=a/b if b and a and b>0 and a>0 else None
    y=replication.get(seq,{}); ys=y.get('values',(None,)*5)
    rb,ra=ys[0],ys[2];rfc=ra/rb if rb and ra and rb>0 and ra>0 else None
    all_rows.append(dict(peptide=seq,source_genes=gene,reference_gene=gene_ref,uniprot=uniprot,start=start,end=end,canonical_mapping_valid=mapping,construct_region=segment,predicted_hla=hla,netmhcpan_rank=rank,basal_identifications=int(str(basal).split('/')[0]),ifna_identifications=int(str(ifna).split('/')[0]),n_biological_replicates=4,mean_basal_intensity=b,mean_ifna_intensity=a,ifna_fold_change=fc,log2_fold_change=math.log2(fc) if fc else None,independent_median_basal=rb,independent_median_ifna=ra,independent_fold_change=rfc,independent_n_biological_replicates=3,table2_source_cell=f'cleanpep!A{rownum}',quant_source_cell=x.get('source_cell'),independent_source_cell=y.get('source_cell')))
assert len(all_rows)==784
assert len({r['peptide'] for r in all_rows})==784
assert all(r['canonical_mapping_valid'] for r in all_rows if r['reference_gene']=='INS')
write_csv('carre2025_conventional_joined.csv',all_rows)
selected=[]
selection=['ALWGPDPAAA','HLVEALYLV','GERGFFYTP','RSAAVLALL','VMNILLQYV','VMNILLQYVV','VAANIVLTV','KMFPEVKEK','KVAPVIKAR','RTYLVNDKAAK']
byseq={r['peptide']:r for r in all_rows}
for seq in selection:
    if seq in byseq:
        r=dict(byseq[seq]);r['status']='detected';selected.append(r)
    else:
        r={k:None for k in all_rows[0]};r['peptide']=seq;r['status']='not present in conventional table; not a zero';selected.append(r)
write_csv('selected_targets.csv',selected)
ins=[r for r in all_rows if r['reference_gene']=='INS']
write_csv('ins_canonical_projection.csv',ins)

def count_region(rows):
    return dict(total_distinct_sequences=len(rows),basal_any=sum(r['basal_identifications']>0 for r in rows),ifna_any=sum(r['ifna_identifications']>0 for r in rows),basal_ge2=sum(r['basal_identifications']>=2 for r in rows),ifna_ge2=sum(r['ifna_identifications']>=2 for r in rows),both_4of4=sum(r['basal_identifications']==4 and r['ifna_identifications']==4 for r in rows),quantified_both=sum(r['ifna_fold_change'] is not None for r in rows),median_fold_quantified=median([r['ifna_fold_change'] for r in rows if r['ifna_fold_change']]) if any(r['ifna_fold_change'] for r in rows) else None)
counts={reg:count_region([r for r in ins if r['construct_region']==reg]) for reg in sorted(set(r['construct_region'] for r in ins))}
sp=byseq['ALWGPDPAAA'];b10=byseq['HLVEALYLV']
ins_quant=[r for r in ins if r['ifna_fold_change'] is not None and r['basal_identifications']>=2 and r['ifna_identifications']>=2]
ins_quant.sort(key=lambda r:r['ifna_fold_change'],reverse=True)
for i,r in enumerate(ins_quant,1):r['fold_induction_rank_among_ins_detected_ge2_each_condition']=i
write_csv('ins_robust_detected_induction.csv',ins_quant)
summary=dict(analysis='Canonical INS signal-peptide replacement projection on wild-type Carré2025 ECN90; not a measured edited-cell outcome',source_url=SRC,pmid='39824805',source_data_url='https://static-content.springer-cdn.com/esm/art%3A10.1038%2Fs41467-025-55908-9/MediaObjects/41467_2025_55908_MOESM14_ESM.xlsx',quantitative_method='Main: source-reported mean total Progenesis ion abundance across 4 biological replicates, Fusion Lumos. Independent: source-reported median peptide intensity summary across a 3-biological-replicate proteasome experiment, timsTOF SCP; compare basal and IFNα-only arms. No per-replicate B10–18 or CHGA2–10 intensities recovered; Fig8c supplies separate run-level values for SP15–24 and GER44–52.',quantified_unique_sequences=len(quant),conventional_sequences=784,ins_counts=counts,b10_relative_to_sp_induction=b10['ifna_fold_change']/sp['ifna_fold_change'],independent_b10_relative_to_sp_induction=b10['independent_fold_change']/sp['independent_fold_change'],hla_assignment_note='SuppData2 NetMHCpan4.1a predictions. RSAAVLALL E01:01 here, not experimentally proven allele-specific source of its measured increase.',sequence_mapping_failures=[dict(peptide=r['peptide'],gene=r['reference_gene'],reported_start=r['start'],reported_end=r['end']) for r in all_rows if r['canonical_mapping_valid'] is False],input_sha256={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in ROOT.glob('*.xlsx')},limitations=['No replicate-level variance or target-specific p-values for B10–18 or CHGA2–10 can be calculated from condition summaries. Fig8c run-level values for SP and GER are separately audited.','Peptide ion response factors differ; fold changes within a peptide are valid summaries, but raw intensities across different peptides do not estimate molecular abundance ratios.','Whole INS coordinates1-24 are a canonical projection. Exact replacement and junction sequence not available. Sequence retention does not establish unchanged processing in the engineered cells.','Candidate lists are beta-cell-enriched peptide sequences rather than all MS peptides. Shared INS/INS-IGF2 source mapping is preserved.','Main and independent measurements have different platforms and aggregations; magnitude does not replicate closely.','MS non-detection is not absence. Classical-A/B/C vs HLA-E restriction is not resolved by pan-HLA immunoprecipitation.'])
(ROOT/'results.json').write_text(json.dumps(summary,indent=2))
print(json.dumps({k:v for k,v in summary.items() if k not in ['input_sha256','limitations']},indent=2))
print('SELECTED:',[(r['peptide'],r['ifna_fold_change'],r['independent_fold_change']) for r in selected])
print('ROBUST INS:',[(r['peptide'],r['ifna_fold_change']) for r in ins_quant])

# Independent check of Figure 8c run-level abundances against Figure 4 summaries.
f8=sheet('41467_2025_55908_MOESM14_ESM.xlsx','Fig 8a-b-c-d')
f8rows=[]
qc=[]
for seq,begin in [('ALWGPDPAAA',20),('GERGFFYTP',24)]:
    for i,r in enumerate(f8[begin:begin+4],1):
        f8rows.append(dict(peptide=seq,ms_run=i,basal_intensity=r[5],ifna_intensity=r[6],source_cell=f'Fig 8a-b-c-d!F{begin+i}'))
    z=[r for r in f8rows if r['peptide']==seq]
    mb=sum(r['basal_intensity'] for r in z)/4;ma=sum(r['ifna_intensity'] for r in z)/4
    q=dict(peptide=seq,figure8c_mean_basal=mb,figure8c_mean_ifna=ma,figure8c_fold=ma/mb,figure4_mean_basal=byseq[seq]['mean_basal_intensity'],figure4_mean_ifna=byseq[seq]['mean_ifna_intensity'],figure4_fold=byseq[seq]['ifna_fold_change'])
    q['figure4_over_figure8_basal']=q['figure4_mean_basal']/mb
    q['figure4_over_figure8_ifna']=q['figure4_mean_ifna']/ma
    qc.append(q)
write_csv('figure8c_reported_run_intensities.csv',f8rows)
write_csv('source_summary_consistency.csv',qc)
summary['figure4_figure8_consistency']=qc
summary['limitations'].append('Figure8c SP15–24 run-level means give 0.904-fold IFN/basal, versus1.089-fold inFigure4; GER44–52 means reconcile. The reason for the SP discrepancy is not specified. The exact10.33x B10/SP differential is Figure4-specific.')
(ROOT/'results.json').write_text(json.dumps(summary,indent=2))
