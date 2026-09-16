"""Build an editable journal manuscript from the canonical Markdown source.

Requires python-docx. Figures and numeric results are generated separately from
cached data. This builder only formats the manuscript and verified references.
"""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
REF_ORDER = ['kobaisi2026', 'whalley2020', 'nanaware2025', 'carre2025',
             'reynisson2020', 'odonnell2020', 'teufel2022', 'schlosser2007',
             'uniprot2025', 'iedb2025', 'iedbtools2024']


def references():
    source = json.loads((ROOT / 'references.json').read_text())
    records = {r['id']: r for r in source['reference_records']}
    paragraphs = []
    for i, key in enumerate(REF_ORDER, 1):
        r = records[key]
        names = [a.get('indexed_name') or a.get('family') or a.get('collective_name') for a in r['authors']]
        names = [n for n in names if n]
        if len(names) > 10:
            names = names[:10] + ['et al.']
        authors = ', '.join(names)
        if not authors:
            authors = 'UniProt Consortium' if key == 'uniprot2025' else ''
        assert authors, key
        title = r['title'].rstrip('.')
        text = (f"{i}. {authors} ({r['year']}). {title}. "
                f"{r['journal_abbreviation']} {r['volume']}, {r['pages_or_article_number']}. "
                f"[doi:{r['doi']}]({r['url']})")
        paragraphs.append(text)
    (ROOT / 'references.bib').write_text('\n\n'.join(
        '@article{' + key + ',\n' +
        f"  title = {{{records[key]['title'].rstrip('.')}}},\n" +
        '  author = {' + ' and '.join(
            (a.get('family', '') + ', ' + a.get('given', '')).strip(', ')
            for a in records[key]['authors']) + '},\n' +
        f"  journal = {{{records[key]['journal']}}},\n" +
        f"  year = {{{records[key]['year']}}},\n" +
        f"  volume = {{{records[key]['volume']}}},\n" +
        f"  pages = {{{records[key]['pages_or_article_number']}}},\n" +
        f"  doi = {{{records[key]['doi']}}}\n" + '}' for key in REF_ORDER))
    return '\n\n'.join(paragraphs)


def hyperlink(paragraph, label, url):
    relationship = paragraph.part.relate_to(
        url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), relationship)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '174C72'); props.append(color)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def inline(paragraph, text):
    text = text.replace('\\*', '*').replace('\\_', '_')
    pattern = re.compile(r'(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)')
    for chunk in pattern.split(text):
        if chunk.startswith('**') and chunk.endswith('**'):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith('[') and '](' in chunk:
            label, url = chunk[1:].split('](', 1)
            hyperlink(paragraph, label, url[:-1])
        elif chunk.startswith('`') and chunk.endswith('`'):
            run = paragraph.add_run(chunk[1:-1]); run.font.name = 'Courier New'
            run.font.size = Pt(9)
        else:
            paragraph.add_run(chunk)


def table(doc, lines, number):
    rows = [[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', c) for c in r)]
    widths = {1: [1.05, 1.40, 1.05, 1.05, 1.05, 1.30],
              2: [3.30, 1.35, 1.35, .90], 3: [1.75, 2.25, 2.90]}[number]
    assert len(widths) == len(rows[0])
    tbl = doc.add_table(rows=0, cols=len(rows[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for col, width in zip(tbl.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:' + side); el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
    tbl._tbl.tblPr.append(borders)
    for n, values in enumerate(rows):
        row = tbl.add_row()
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if n == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for c, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[c]; cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margin = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '90'); el.set(qn('w:type'), 'dxa'); margin.append(el)
            pr.append(margin)
            if n == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9EDF0'); pr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = True
            if c >= (2 if number == 1 else 1) and number != 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value)
            for run in p.runs:
                run.font.name = 'Arial'; run.font.size = Pt(8.5 if number == 1 else 9)
                run.bold = n == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.keep_with_next = True


def build():
    text = (ROOT / 'manuscript.md').read_text()
    text = text.replace('<!-- REFERENCES_INSERT -->', references())
    captions = ROOT / 'figures/captions.md'
    if captions.exists():
        text = text.replace('<!-- FIGURES_INSERT -->', captions.read_text())
    else:
        text = text.replace('<!-- FIGURES_INSERT -->', '')
    assert '\u2014' not in text
    (ROOT / 'article.md').write_text(text)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.75)
    sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
    sec.header_distance = Inches(.3); sec.footer_distance = Inches(.3)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name, size in [('Title', 19), ('Heading 1', 13), ('Heading 2', 11.5), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = 'Arial'; style.font.size = Pt(size)
        style.font.bold = True; style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].paragraph_format.space_after = Pt(14)
    doc.styles['Caption'].font.name = 'Times New Roman'
    doc.styles['Caption'].font.size = Pt(10)
    doc.styles['Caption'].font.color.rgb = RGBColor(0,0,0)
    doc.styles['Caption'].font.bold = False
    doc.styles['Caption'].paragraph_format.keep_with_next = False
    foot = sec.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    doc.core_properties.title = text.splitlines()[0][2:]
    doc.core_properties.subject = 'Computational research on engineered insulin signal peptide variants'
    doc.core_properties.author = 'GloGlo Labs'
    doc.core_properties.keywords = 'insulin; chromogranin A; HLA; computational antigen engineering'
    lines = []
    pending = []
    for line in text.splitlines():
        if not line.strip() or line.startswith(('#', '|', '- ', '![', '<!--')):
            if pending:
                lines.append(' '.join(pending)); pending = []
            lines.append(line)
        else:
            pending.append(line.strip())
    if pending: lines.append(' '.join(pending))
    i = 0; table_n = 0; in_refs = False
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith('<!--'):
            i += 1; continue
        if line.startswith('|'):
            group = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                group.append(lines[i]); i += 1
            table_n += 1; table(doc, group, table_n); continue
        if line.startswith('# '):
            p = doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            label = line[3:]
            in_refs = label == 'References'
            if label == 'Figures':
                doc.add_page_break()
            if label == 'Figures':
                i += 1; continue
            p = doc.add_paragraph(label, 'Heading 1')
        elif line.startswith('### '):
            label = line[4:]
            if label.startswith('Figure '):
                if not label.startswith('Figure 1 '): doc.add_page_break()
            p = doc.add_paragraph(label, 'Heading 2')
        elif line.startswith('!['):
            match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
            assert match, line
            alt, filename = match.groups()
            path = (ROOT / filename).resolve()
            assert path.is_file(), path
            p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
            run = p.add_run(); picture = run.add_picture(str(path), width=Inches(6.9))
            picture._inline.docPr.set('descr', alt)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); inline(p, line[2:])
            p.paragraph_format.space_after = Pt(3)
        else:
            if line.startswith('Table '):
                p = doc.add_paragraph(style='Caption')
            else:
                p = doc.add_paragraph()
            inline(p, line)
            if in_refs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs: run.font.size = Pt(10)
        i += 1
    output = ROOT / 'chga-ins-leader-variants.docx'
    doc.save(output)
    print(json.dumps({'docx':str(output),'words':len(text.split()),'tables':table_n,'figures':len(doc.inline_shapes)}))


if __name__ == '__main__':
    build()
